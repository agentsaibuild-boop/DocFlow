from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from docflow.schema import ExtractedDocument, InvoiceData


def write_excel(doc: ExtractedDocument, output_path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    invoice_has_line_items = bool(doc.invoice and doc.invoice.line_items)

    if doc.invoice is not None:
        _write_invoice_sheets(wb, doc.invoice)

    if doc.tables and not invoice_has_line_items:
        for idx, table in enumerate(doc.tables, start=1):
            sheet_name = f"p{table.page}_t{idx}"[:31]
            sheet = wb.create_sheet(sheet_name)
            for row in table.rows:
                sheet.append(row)

    if doc.full_text and not doc.tables and doc.invoice is None:
        sheet = wb.create_sheet("text")
        sheet["A1"] = "Full text"
        sheet["A2"] = doc.full_text

    meta = wb.create_sheet("_meta")
    meta.append(["source_path", doc.source_path])
    meta.append(["extraction_method", doc.extraction_method])
    meta.append(["extracted_at", doc.extracted_at.isoformat()])
    meta.append(["page_count", doc.page_count])
    meta.append(["table_count", len(doc.tables)])

    wb.save(output_path)


def _write_invoice_sheets(wb: Workbook, inv: InvoiceData) -> None:
    bold = Font(bold=True)
    s = inv.supplier or type(inv.supplier)() if inv.supplier is not None else None
    c = inv.customer or type(inv.customer)() if inv.customer is not None else None
    from docflow.schema import Party
    s = inv.supplier if inv.supplier is not None else Party()
    c = inv.customer if inv.customer is not None else Party()

    header = wb.create_sheet("invoice")
    rows = [
        ("Поле", "Стойност"),
        ("Номер", inv.invoice_number),
        ("Дата издаване", inv.issue_date),
        ("Дата на предоставяне", inv.delivery_date),
        ("Срок на плащане", inv.payment_due_date),
        ("Валута", inv.currency or "EUR"),
        ("", ""),
        ("--- ИЗПЪЛНИТЕЛ ---", ""),
        ("Име", s.name),
        ("Адрес", s.address),
        ("ЕИК", s.eik),
        ("ИН по ДДС", s.vat_number),
        ("МОЛ", s.mol),
        ("Телефон", s.phone),
        ("Email", s.email),
        ("", ""),
        ("--- ПОЛУЧАТЕЛ ---", ""),
        ("Име", c.name),
        ("Адрес", c.address),
        ("ЕИК", c.eik),
        ("ИН по ДДС", c.vat_number),
        ("МОЛ", c.mol),
        ("", ""),
        ("--- ПЛАЩАНЕ ---", ""),
        ("Метод", inv.payment_method),
        ("IBAN", inv.iban),
        ("Банка", inv.bank),
        ("BIC", inv.bic),
        ("", ""),
        ("--- СУМИ ---", ""),
        ("Сума без отстъпка", inv.subtotal),
        ("Отстъпка", inv.discount),
        ("Обща нетна сума", inv.net_amount),
        ("За плащане", inv.total_to_pay),
        ("Платени", inv.paid),
        ("Остава", inv.remaining),
    ]
    for r in rows:
        header.append(r)
    header["A1"].font = bold
    header["B1"].font = bold
    header.column_dimensions["A"].width = 28
    header.column_dimensions["B"].width = 40

    if inv.line_items:
        items = wb.create_sheet("line_items")
        items.append([
            "№", "Описание", "К-во", "ME", "Ед. цена",
            "Отстъпка %", "Цена с отст.", "ДДС %", "ДДС", "Общо без ДДС",
        ])
        for cell in items[1]:
            cell.font = bold
        for li in inv.line_items:
            items.append([
                li.number,
                li.description,
                li.quantity,
                li.unit,
                li.unit_price,
                li.discount_percent,
                li.price_after_discount,
                li.vat_percent,
                li.vat_amount,
                li.total_without_vat,
            ])

    if inv.vat_breakdown:
        vat = wb.create_sheet("vat")
        vat.append(["ДДС %", "Данъчна основа", "ДДС"])
        for cell in vat[1]:
            cell.font = bold
        for v in inv.vat_breakdown:
            vat.append([v.rate_percent, v.base_amount, v.vat_amount])

    if inv.notes:
        notes = wb.create_sheet("notes")
        notes["A1"] = "Бележки"
        notes["A1"].font = bold
        notes["A2"] = inv.notes
