from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from docflow.pipeline import NoExtractorFound, ProviderError, extract
from docflow.registry import SupplierRegistry
from docflow.schema import ExtractedDocument
from docflow.status import compute_status
from docflow.validators import validate

SUPPORTED = {".pdf", ".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp", ".webp"}


@dataclass
class BatchResult:
    source: Path
    output: Path | None
    doc: ExtractedDocument | None
    error: str | None
    enrich_findings: list
    validation_findings: list
    registry_findings: list
    provider_error: str | None = None
    provider_error_kind: str | None = None  # quota | auth | network | other
    processing_time_s: float | None = None  # wall-clock for extract+validate+registry


def discover(folder: Path) -> list[Path]:
    return sorted(p for p in folder.rglob("*") if p.is_file() and p.suffix.lower() in SUPPORTED)


def process_one(path: Path, registry: SupplierRegistry, provider: str = "auto",
                allow_fallback: bool | None = None) -> BatchResult:
    import time
    t0 = time.time()
    try:
        doc = extract(path, provider=provider, allow_fallback=allow_fallback)
    except ProviderError as pe:
        return BatchResult(
            path, None, None, None, [], [], [],
            provider_error=str(pe), provider_error_kind=pe.kind,
            processing_time_s=round(time.time() - t0, 2),
        )
    except Exception as e:
        return BatchResult(
            path, None, None, f"{type(e).__name__}: {e}", [], [], [],
            processing_time_s=round(time.time() - t0, 2),
        )

    enrich_findings = []
    if doc.invoice:
        doc.invoice, enrich_findings = registry.enrich(doc.invoice)

    validation_findings = validate(doc)
    registry_findings = registry.record(doc.invoice, human_confirmed=False) if doc.invoice else []

    from docflow.quality import invoice_quality_coverage, invoice_quality_score
    doc.quality_score = invoice_quality_score(doc.invoice, validation_findings)
    doc.quality_coverage = invoice_quality_coverage(doc.invoice)

    return BatchResult(
        path, None, doc, None,
        enrich_findings, validation_findings, registry_findings,
        processing_time_s=round(time.time() - t0, 2),
    )


def _provider_name_from_method(method: str) -> str:
    """Strip model variant suffix to get the underlying provider/extractor name.

    'gemini:gemini-3.1-flash-lite' → 'gemini'
    'pdfplumber+gemini_text'       → 'pdfplumber'
    'qwen_3_vl_235b:qwen3-vl-235b' → 'qwen_3_vl_235b'
    """
    base = method.split("+")[0]
    base = base.split(":")[0]
    return base


def summarize_batch(results: list[BatchResult]) -> dict:
    """Aggregate batch counts. Provider failures are kept separate from
    document quality. Per-provider stats included for inspection.

    Top-level (mutually exclusive except 'processed'):
      processed, extracted_ok, validation_errors, no_data,
      provider_failures, unknown_errors

    per_provider: {provider_name: {files, avg_score, avg_coverage,
                                   validation_errors, avg_latency_s}}
    Only successful extractions contribute to per_provider averages; provider
    failures are counted globally.
    """
    processed = len(results)
    provider_failures = sum(1 for r in results if r.provider_error)
    unknown_errors = sum(1 for r in results if r.error)
    validation_errors = 0
    extracted_ok = 0
    no_data = 0

    per_provider: dict[str, dict] = {}

    for r in results:
        if r.provider_error or r.error:
            continue
        if not (r.doc and r.doc.invoice and r.doc.invoice.supplier and r.doc.invoice.supplier.name):
            no_data += 1
            continue

        has_validation_error = any(getattr(f, "level", None) == "error" for f in r.validation_findings)
        if has_validation_error:
            validation_errors += 1
        else:
            extracted_ok += 1

        provider = _provider_name_from_method(r.doc.extraction_method)
        s = per_provider.setdefault(provider, {
            "files": 0, "_score_sum": 0.0, "_coverage_sum": 0.0,
            "validation_errors": 0, "_latency_sum": 0.0, "_latency_count": 0,
        })
        s["files"] += 1
        s["_score_sum"] += r.doc.quality_score
        s["_coverage_sum"] += getattr(r.doc, "quality_coverage", 0.0)
        if has_validation_error:
            s["validation_errors"] += 1
        if r.processing_time_s is not None:
            s["_latency_sum"] += r.processing_time_s
            s["_latency_count"] += 1

    # Finalize per-provider averages, drop internal accumulators.
    for provider, s in per_provider.items():
        n = s["files"]
        s["avg_score"]    = round(s["_score_sum"] / n, 3) if n else 0.0
        s["avg_coverage"] = round(s["_coverage_sum"] / n, 3) if n else 0.0
        s["avg_latency_s"] = round(s["_latency_sum"] / s["_latency_count"], 2) if s["_latency_count"] else None
        for k in ("_score_sum", "_coverage_sum", "_latency_sum", "_latency_count"):
            del s[k]

    return {
        "processed": processed,
        "extracted_ok": extracted_ok,
        "validation_errors": validation_errors,
        "no_data": no_data,
        "provider_failures": provider_failures,
        "unknown_errors": unknown_errors,
        "per_provider": per_provider,
    }


def run_batch(folder: Path, output_dir: Path, provider: str = "auto",
              allow_fallback: bool | None = None) -> list[BatchResult]:
    output_dir.mkdir(parents=True, exist_ok=True)
    files = discover(folder)
    if not files:
        return []
    registry = SupplierRegistry()
    results = []
    for f in files:
        print(f"[{len(results) + 1}/{len(files)}] {f.relative_to(folder)}", flush=True)
        results.append(process_one(f, registry, provider=provider, allow_fallback=allow_fallback))
    write_consolidated(results, output_dir / "Фактури.xlsx")
    summary = summarize_batch(results)
    print(
        "  done: processed={processed} extracted_ok={extracted_ok} "
        "validation_errors={validation_errors} no_data={no_data} "
        "provider_failures={provider_failures} unknown_errors={unknown_errors}".format(**summary),
        flush=True,
    )
    return results


def write_consolidated(results: list[BatchResult], output_path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    bold = Font(bold=True)

    inv_sheet = wb.create_sheet("Фактури")
    inv_headers = [
        "Файл", "Статус", "Метод",
        "Доставчик", "ЕИК доставчик", "ВАТ доставчик",
        "Получател", "ЕИК получател", "ВАТ получател",
        "Номер", "Дата издаване", "Дата доставка", "Срок плащане",
        "Валута",
        "Сума без отстъпка", "Отстъпка", "Нетна сума",
        "За плащане", "Платени", "Остава",
        "IBAN", "Банка", "BIC", "Метод плащане",
        "Грешка",
    ]
    inv_sheet.append(inv_headers)
    for cell in inv_sheet[1]:
        cell.font = bold

    items_sheet = wb.create_sheet("Артикули")
    items_sheet.append([
        "Файл", "Доставчик", "Номер фактура",
        "№", "Описание", "К-во", "ME",
        "Ед. цена", "Отстъпка %", "Цена с отст.",
        "ДДС %", "ДДС", "Общо без ДДС",
    ])
    for cell in items_sheet[1]:
        cell.font = bold

    vat_sheet = wb.create_sheet("ДДС")
    vat_sheet.append(["Файл", "Доставчик", "Номер фактура", "Ставка %", "Основа", "ДДС"])
    for cell in vat_sheet[1]:
        cell.font = bold

    val_sheet = wb.create_sheet("Validation")
    val_sheet.append(["Файл", "Източник", "Ниво", "Код", "Съобщение"])
    for cell in val_sheet[1]:
        cell.font = bold

    for r in results:
        if r.error:
            status_label = compute_status(None, [], extraction_error=r.error).label
            inv_sheet.append([r.source.name, status_label, "", *[""] * 21, r.error])
            continue
        doc = r.doc
        inv = doc.invoice
        status_label = compute_status(doc, r.validation_findings).label

        if inv is None:
            inv_sheet.append([r.source.name, status_label, doc.extraction_method, *[""] * 22])
        else:
            inv_sheet.append([
                r.source.name, status_label, doc.extraction_method,
                inv.supplier.name, inv.supplier.eik, inv.supplier.vat_number,
                inv.customer.name, inv.customer.eik, inv.customer.vat_number,
                inv.invoice_number, inv.issue_date, inv.delivery_date, inv.payment_due_date,
                inv.currency or "EUR",
                inv.subtotal, inv.discount, inv.net_amount,
                inv.total_to_pay, inv.paid, inv.remaining,
                inv.iban, inv.bank, inv.bic, inv.payment_method,
                "",
            ])

            for li in (inv.line_items or []):
                items_sheet.append([
                    r.source.name, inv.supplier.name, inv.invoice_number,
                    li.number, li.description, li.quantity, li.unit,
                    li.unit_price, li.discount_percent, li.price_after_discount,
                    li.vat_percent, li.vat_amount, li.total_without_vat,
                ])

            for v in (inv.vat_breakdown or []):
                vat_sheet.append([
                    r.source.name, inv.supplier.name, inv.invoice_number,
                    v.rate_percent, v.base_amount, v.vat_amount,
                ])

        for f in r.enrich_findings:
            val_sheet.append([r.source.name, "registry_enrich", f.level, f.code, f.message])
        for f in r.validation_findings:
            val_sheet.append([r.source.name, "validation", f.level, f.code, f.message])
        for f in r.registry_findings:
            val_sheet.append([r.source.name, "registry_record", f.level, f.code, f.message])

    meta_sheet = wb.create_sheet("_meta")
    meta_sheet.append(["batch_run_at", datetime.now().isoformat()])
    meta_sheet.append(["files_total", len(results)])
    meta_sheet.append(["files_ok", sum(1 for r in results if r.error is None)])
    meta_sheet.append(["files_error", sum(1 for r in results if r.error is not None)])
    method_counts = {}
    for r in results:
        if r.doc:
            method_counts[r.doc.extraction_method] = method_counts.get(r.doc.extraction_method, 0) + 1
    for m, c in sorted(method_counts.items(), key=lambda x: -x[1]):
        meta_sheet.append([f"метод: {m}", c])

    inv_sheet.freeze_panes = "A2"
    items_sheet.freeze_panes = "A2"
    vat_sheet.freeze_panes = "A2"
    val_sheet.freeze_panes = "A2"

    inv_sheet.column_dimensions["A"].width = 35
    inv_sheet.column_dimensions["D"].width = 30
    inv_sheet.column_dimensions["G"].width = 30
    items_sheet.column_dimensions["A"].width = 35
    items_sheet.column_dimensions["B"].width = 25
    items_sheet.column_dimensions["E"].width = 35
    val_sheet.column_dimensions["A"].width = 35
    val_sheet.column_dimensions["E"].width = 60

    wb.save(output_path)
